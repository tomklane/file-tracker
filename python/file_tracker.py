#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
File Tracker
------------
Scans a folder (as seen in Finder / File Explorer) and fills a tracking
spreadsheet (.xlsx, .xlsm or .csv) using whatever columns that spreadsheet
already has. You choose what goes in each column; columns you leave alone
(status, notes, reviewer...) are never touched.

Runs on macOS and Windows with Python 3.8+.
Needs the free 'openpyxl' package, which it installs by itself on first run.
"""

import csv
import datetime as dt
import glob
import io
import json
import fnmatch
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
import traceback
from copy import copy
from pathlib import Path

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

APP_NAME = "File Tracker"
SETTINGS_FILE = Path.home() / ".file_tracker_settings.json"
BACKUP_DIR = "_tracker_backups"
MAX_BACKUPS = 20
JUNK_FILES = {".ds_store", "thumbs.db", "desktop.ini", "icon\r", ".localized"}
DEFAULT_NEW_COLUMNS = ("ID, File Name, Type, Folder, Full Path, Size, "
                       "Date Modified, Date Added, On Disk, Status, Notes")
CSV_EXTS = (".csv", ".tsv", ".txt")
XLSX_EXTS = (".xlsx", ".xlsm")

# key, label shown in the app, default setting (None = no setting), hint for the setting box
FIELDS = [
    ("",           "— leave as is —",                      None, ""),
    ("name",       "File name",                            None, ""),
    ("stem",       "File name without extension",          None, ""),
    ("part",       "Part of the file name",                "1",  "1 = first part, -1 = last"),
    ("ext",        "Extension (pdf)",                      None, ""),
    ("ext_upper",  "Type (PDF)",                           None, ""),
    ("parent",     "Folder it's in",                       None, ""),
    ("level",      "Folder at level…",                     "1",  "1 = top subfolder, -1 = its own folder"),
    ("rel_folder", "Folder path inside scanned folder",    None, ""),
    ("rel_path",   "File path inside scanned folder",      None, ""),
    ("full_path",  "Full file path",                       None, ""),
    ("link",       "Clickable link to the file",           None, ""),
    ("size_h",     "Size (e.g. 2.4 MB)",                   None, ""),
    ("size_kb",    "Size in KB",                           None, ""),
    ("size_mb",    "Size in MB",                           None, ""),
    ("size_b",     "Size in bytes",                        None, ""),
    ("modified_d", "Date modified",                        None, ""),
    ("modified",   "Date & time modified",                 None, ""),
    ("created_d",  "Date created",                         None, ""),
    ("created",    "Date & time created",                  None, ""),
    ("scanned",    "Date added to sheet (new rows only)",  None, ""),
    ("index",      "Running number / ID (new rows only)",  "#",  "e.g.  #  or  DOC-###"),
    ("status",     "Present / Missing",                    "Present|Missing", "found text | missing text"),
    ("text",       "Fixed text",                           "",   "text to write"),
]
FIELD_BY_KEY = {f[0]: f for f in FIELDS}
LABEL_TO_KEY = {f[1]: f[0] for f in FIELDS}
FIELD_LABELS = [f[1] for f in FIELDS]
NEW_ONLY = {"scanned", "index"}                 # written once, when a row is added
BAD_KEYS = {"", "scanned", "index", "status", "text"}
KEY_PRIORITY = ["full_path", "rel_path", "link", "name", "stem"]


class UserError(Exception):
    """A problem we can explain to the user in plain words."""


# --------------------------------------------------------------------------- setup

def ensure_openpyxl():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        pass
    flags = 0x08000000 if os.name == "nt" else 0  # CREATE_NO_WINDOW
    attempts = [
        [sys.executable, "-m", "pip", "install", "--user", "openpyxl"],
        [sys.executable, "-m", "pip", "install", "openpyxl"],
        [sys.executable, "-m", "pip", "install", "--user", "--break-system-packages", "openpyxl"],
    ]
    for cmd in attempts:
        try:
            subprocess.run(cmd, check=True, capture_output=True, creationflags=flags)
        except Exception:
            continue
        try:
            import importlib
            import site
            importlib.invalidate_caches()
            user_site = site.getusersitepackages()
            if user_site not in sys.path:
                sys.path.append(user_site)
            import openpyxl  # noqa: F401
            return True
        except Exception:
            continue
    return False


# --------------------------------------------------------------------------- helpers

def clean(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d")
    return " ".join(str(v).split())


def dedupe_headers(pairs):
    """[(col, text)] -> same, with duplicate header names made unique."""
    seen, out = {}, []
    for col, text in pairs:
        if not text:
            continue
        k = text.casefold()
        seen[k] = seen.get(k, 0) + 1
        out.append((col, text if seen[k] == 1 else f"{text} ({seen[k]})"))
    return out


def detect_header_row(rows):
    """Pick the row (1-based) with the most text cells in the first 30 rows."""
    best, best_n = 1, 0
    for i, row in enumerate(rows[:30], start=1):
        n = sum(1 for v in row if isinstance(v, str) and v.strip())
        if n > best_n:
            best, best_n = i, n
    return best


def to_int(s, default):
    try:
        return int(str(s).strip())
    except (TypeError, ValueError):
        return default


def human_size(n):
    for unit in ("bytes", "KB", "MB", "GB", "TB"):
        if n < 1024 or unit == "TB":
            return f"{n:.0f} {unit}" if unit == "bytes" else f"{n:.1f} {unit}"
        n /= 1024.0


def status_texts(param):
    parts = (param or "Present|Missing").split("|")
    found = parts[0].strip() or "Present"
    missing = parts[1].strip() if len(parts) > 1 and parts[1].strip() else "Missing"
    return found, missing


def norm_key(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, dt.datetime):
        v = v.strftime("%Y-%m-%d %H:%M")
    elif isinstance(v, dt.date):
        v = v.isoformat()
    return str(v).strip().replace("\\", "/").casefold()


def key_text(field, value, is_csv):
    if field == "link":
        value = value[2] if is_csv else value[0]
    return norm_key(value)


def split_list(s):
    return [x.strip() for x in re.split(r"[,;]", s or "") if x.strip()]


def open_path(p):
    p = str(p)
    if sys.platform == "win32":
        os.startfile(p)  # noqa
    elif sys.platform == "darwin":
        subprocess.run(["open", p])
    else:
        subprocess.run(["xdg-open", p])


def auto_field(header):
    """Guess a sensible field from a column header. '' means leave it alone."""
    t = re.sub(r"[^a-z0-9#]+", " ", header.casefold()).strip()
    words = set(t.split())
    has = lambda *ps: any(p in t for p in ps)  # noqa: E731
    if words & {"by", "owner", "author", "reviewer", "comment", "comments", "notes",
                "remarks", "status", "description", "approved", "approver"}:
        return ""
    if t in ("#", "no", "nr", "item", "row", "s no", "sl no", "sr no", "index", "id"):
        return "index"
    if has("hyperlink", "link"):
        return "link"
    if has("on disk", "exists", "file found", "present"):
        return "status"
    if has("relative path", "path in", "subpath"):
        return "rel_path"
    if has("folder path", "directory path"):
        return "rel_folder"
    if has("full path", "file path", "filepath", "location") or t == "path":
        return "full_path"
    if has("folder", "directory"):
        return "parent"
    if has("extension") or "ext" in words:
        return "ext"
    if has("file type", "filetype", "format") or t == "type":
        return "ext_upper"
    if "size" in words:
        if "mb" in words:
            return "size_mb"
        if "kb" in words:
            return "size_kb"
        if words & {"bytes", "byte", "b"}:
            return "size_b"
        return "size_h"
    if has("created", "creation"):
        return "created_d"
    if has("modified", "last updated", "updated", "last changed", "last saved"):
        return "modified_d"
    if has("date added", "added", "logged", "date entered", "first seen"):
        return "scanned"
    if has("file name", "filename") or t in ("name", "file"):
        return "name"
    if has("title", "document name", "doc name"):
        return "stem"
    return ""


# --------------------------------------------------------------------------- scanning

def is_hidden(p, st=None):
    if p.name.startswith("."):
        return True
    try:
        st = st or p.stat()
    except OSError:
        return False
    if os.name == "nt":
        return bool(getattr(st, "st_file_attributes", 0) & (0x2 | 0x4))  # hidden | system
    if sys.platform == "darwin":
        return bool(getattr(st, "st_flags", 0) & 0x8000)  # UF_HIDDEN
    return False


def matches_any(name, patterns):
    n = name.casefold()
    return any(fnmatch.fnmatch(n, pat.casefold()) for pat in patterns)


def created_time(st):
    ts = getattr(st, "st_birthtime", None)
    if ts is None and os.name == "nt":
        ts = st.st_ctime
    return dt.datetime.fromtimestamp(ts).replace(microsecond=0) if ts else None


def scan_folder(root, recursive=True, skip_hidden=True, extensions=None,
                exclude=None, skip_paths=None, progress=None):
    root = Path(root)
    exts = {e.lower().lstrip(".").lstrip("*.") for e in (extensions or [])}
    exclude = list(exclude or []) + [BACKUP_DIR]
    skip_paths = {os.path.normcase(str(p)) for p in (skip_paths or [])}
    out = []
    for dirpath, dirnames, filenames in os.walk(root, onerror=lambda e: None):
        keep = []
        for d in sorted(dirnames, key=str.casefold):
            if matches_any(d, exclude):
                continue
            if skip_hidden and is_hidden(Path(dirpath, d)):
                continue
            keep.append(d)
        dirnames[:] = keep if recursive else []
        for fn in sorted(filenames, key=str.casefold):
            low = fn.casefold()
            if low in JUNK_FILES or fn.startswith("~$") or fn.startswith("._"):
                continue
            if matches_any(fn, exclude):
                continue
            p = Path(dirpath, fn)
            if exts and p.suffix.lower().lstrip(".") not in exts:
                continue
            try:
                st = p.stat()
            except OSError:
                continue
            if skip_hidden and is_hidden(p, st):
                continue
            if os.path.normcase(str(p.resolve())) in skip_paths:
                continue
            rel = p.relative_to(root)
            out.append({
                "path": p.resolve(),
                "name": p.name,
                "stem": p.stem,
                "ext": p.suffix.lower().lstrip("."),
                "rel": str(rel),
                "rel_folder": str(rel.parent) if str(rel.parent) != "." else "",
                "folders": [root.name] + list(rel.parent.parts),
                "size": st.st_size,
                "mtime": dt.datetime.fromtimestamp(st.st_mtime).replace(microsecond=0),
                "ctime": created_time(st),
            })
            if progress and len(out) % 500 == 0:
                progress(f"Scanning… {len(out):,} files found")
    return out


def field_value(f, param, info, ctx):
    if f == "name":
        return info["name"]
    if f == "stem":
        return info["stem"]
    if f == "part":
        n = to_int(param, 1)
        bits = info["stem"].split(ctx["delimiter"] or "_")
        idx = n - 1 if n > 0 else (n if n < 0 else 0)
        try:
            return bits[idx].strip()
        except IndexError:
            return ""
    if f == "ext":
        return info["ext"]
    if f == "ext_upper":
        return info["ext"].upper()
    if f == "parent":
        return info["folders"][-1]
    if f == "level":
        try:
            return info["folders"][to_int(param, 1)]
        except IndexError:
            return ""
    if f == "rel_folder":
        return info["rel_folder"]
    if f == "rel_path":
        return info["rel"]
    if f == "full_path":
        return str(info["path"])
    if f == "link":
        return (info["name"], info["path"].as_uri(), str(info["path"]))
    if f == "size_h":
        return human_size(info["size"])
    if f == "size_kb":
        return round(info["size"] / 1024, 1)
    if f == "size_mb":
        return round(info["size"] / 1048576, 2)
    if f == "size_b":
        return info["size"]
    if f == "modified":
        return info["mtime"]
    if f == "modified_d":
        return info["mtime"].date()
    if f == "created":
        return info["ctime"] or ""
    if f == "created_d":
        return info["ctime"].date() if info["ctime"] else ""
    if f == "scanned":
        return ctx["today"]
    if f == "status":
        return status_texts(param)[0]
    if f == "text":
        return param or ""
    return ""


# --------------------------------------------------------------------------- spreadsheets

def unsupported(ext):
    if ext == ".xls":
        return UserError("Old .xls files aren't supported.\nOpen it in Excel, choose File › Save As › "
                         "Excel Workbook (.xlsx), then pick the new file.")
    if ext == ".numbers":
        return UserError("Numbers files aren't supported.\nIn Numbers choose File › Export To › Excel, "
                         "then pick the exported file.")
    return UserError(f"This app works with .xlsx, .xlsm and .csv files (not {ext or 'this file'}).")


def read_csv(path):
    raw = Path(path).read_bytes()
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        delim = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
    except csv.Error:
        delim = "\t" if str(path).lower().endswith(".tsv") else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    return rows, {"delimiter": delim}


def load_any(path):
    import openpyxl
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext in XLSX_EXTS:
            return openpyxl.load_workbook(path, keep_vba=(ext == ".xlsm")), {"kind": "xlsx"}
        if ext in CSV_EXTS:
            rows, meta = read_csv(path)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "CSV"
            for r, row in enumerate(rows, start=1):
                for c, v in enumerate(row, start=1):
                    if v != "":
                        ws.cell(row=r, column=c, value=v)
            return wb, {"kind": "csv", **meta}
    except PermissionError:
        raise UserError(f"Couldn't open {path.name}. If it's open in Excel, close it and try again.")
    raise unsupported(ext)


def peek(path, sheet=None, header_row=None):
    """Read just enough of a spreadsheet to list its sheets and column headers."""
    import openpyxl
    path = Path(path)
    ext = path.suffix.lower()
    limit = max(40, header_row or 0)
    if ext in XLSX_EXTS:
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
        except PermissionError:
            raise UserError(f"Couldn't open {path.name}. If it's open in Excel, close it and try again.")
        try:
            names = wb.sheetnames
            if sheet not in names:
                try:
                    sheet = wb.active.title
                except Exception:
                    sheet = names[0]
            ws = wb[sheet]
            rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=limit, values_only=True)]
        finally:
            wb.close()
    elif ext in CSV_EXTS:
        rows, _ = read_csv(path)
        rows = rows[:limit]
        names, sheet = ["(CSV file)"], "(CSV file)"
    else:
        raise unsupported(ext)
    hr = header_row or detect_header_row(rows)
    vals = rows[hr - 1] if 0 < hr <= len(rows) else []
    headers = dedupe_headers([(i + 1, clean(v)) for i, v in enumerate(vals)])
    return {"sheets": names, "sheet": sheet, "header_row": hr, "headers": headers}


def read_headers(ws, hr):
    pairs = [(c, clean(ws.cell(row=hr, column=c).value)) for c in range(1, (ws.max_column or 1) + 1)]
    return dict(dedupe_headers(pairs))


def last_data_row(ws, hr, cols):
    for r in range(ws.max_row, hr, -1):
        if any(ws.cell(row=r, column=c).value not in (None, "") for c in cols):
            return r
    return hr


def new_workbook(columns):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "File Tracker"
    for i, name in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5496")
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    return wb


def put(ws, r, c, f, val, is_csv):
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import Font
    cell = ws.cell(row=r, column=c)
    if isinstance(cell, MergedCell):
        return
    if f == "link":
        name, uri, path = val
        if is_csv:
            cell.value = path
            return
        cell.value = name
        cell.hyperlink = uri
        fnt = cell.font
        cell.font = Font(name=fnt.name, size=fnt.size, bold=fnt.bold, italic=fnt.italic,
                         color="0563C1", underline="single")
        return
    cell.value = val
    if not is_csv:
        if isinstance(val, dt.datetime):
            if cell.number_format in ("General", "yyyy-mm-dd h:mm:ss"):
                cell.number_format = "yyyy-mm-dd hh:mm"
        elif isinstance(val, dt.date):
            if cell.number_format in ("General", "yyyy-mm-dd h:mm:ss"):
                cell.number_format = "yyyy-mm-dd"


def save_any(wb, path, meta):
    path = Path(path)
    if path.suffix.lower() in CSV_EXTS:
        ws = wb.active
        rows = []
        for row in ws.iter_rows(values_only=True):
            out = []
            for v in row:
                if v is None:
                    v = ""
                elif isinstance(v, dt.datetime):
                    v = v.strftime("%Y-%m-%d %H:%M")
                elif isinstance(v, dt.date):
                    v = v.isoformat()
                out.append(v)
            rows.append(out)
        while rows and not any(str(v) for v in rows[-1]):
            rows.pop()
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            csv.writer(fh, delimiter=meta.get("delimiter", ",")).writerows(rows)
    else:
        wb.save(path)


def make_backup(path):
    path = Path(path)
    bdir = path.parent / BACKUP_DIR
    bdir.mkdir(exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y-%m-%d %H%M%S")
    dest = bdir / f"{path.stem} {stamp}{path.suffix}"
    shutil.copy2(path, dest)
    olds = sorted(bdir.glob(f"{glob.escape(path.stem)} *{path.suffix}"), key=os.path.getmtime)
    for old in olds[:-MAX_BACKUPS]:
        try:
            old.unlink()
        except OSError:
            pass
    return dest


def prepare(p, progress=lambda m: None):
    """Open (or create) the sheet and work out which column gets which field."""
    sheet_path = Path(p["sheet_path"]) if p.get("sheet_path") else None
    is_csv = bool(sheet_path) and sheet_path.suffix.lower() in CSV_EXTS
    folder = Path(p["folder"])
    if not folder.is_dir():
        raise UserError("Choose a folder to scan (section 1).")

    if p["create_new"]:
        wb, meta = new_workbook(p["new_columns"]), {"kind": "csv" if is_csv else "xlsx", "delimiter": ","}
        ws, hr = wb.active, 1
    else:
        if not sheet_path or not sheet_path.is_file():
            raise UserError("Choose your tracking spreadsheet (section 2).")
        progress("Opening spreadsheet…")
        wb, meta = load_any(sheet_path)
        if meta["kind"] == "csv":
            ws = wb.active
        else:
            if p["sheet_name"] not in wb.sheetnames:
                raise UserError(f'Sheet "{p["sheet_name"]}" wasn\'t found. Click "Reload columns".')
            ws = wb[p["sheet_name"]]
        hr = int(p["header_row"])

    headers = read_headers(ws, hr)
    if not headers:
        raise UserError(f"No column headers were found on row {hr}.")
    by_text = {}
    for c, t in headers.items():
        by_text.setdefault(t.casefold(), c)

    plan = []
    for m in p["mapping"]:
        if not m["field"]:
            continue
        col = m["col"] if headers.get(m["col"], "").casefold() == m["header"].casefold() \
            else by_text.get(m["header"].casefold())
        if col is None:
            raise UserError(f'Column "{m["header"]}" isn\'t on row {hr} any more. Click "Reload columns" and try again.')
        plan.append((col, m["field"], m["param"]))
    if not plan:
        raise UserError("Choose what goes into at least one column (section 3).")

    key_col = by_text.get((p.get("key_header") or "").casefold())
    key_plan = next((x for x in plan if x[0] == key_col and x[1] not in BAD_KEYS), None)
    if key_plan is None:
        for k in KEY_PRIORITY:
            key_plan = next((x for x in plan if x[1] == k), None)
            if key_plan:
                break
    if key_plan is None:
        key_plan = next((x for x in plan if x[1] not in BAD_KEYS), plan[0])
    return {"wb": wb, "meta": meta, "ws": ws, "hr": hr, "headers": headers, "plan": plan,
            "key_plan": key_plan, "sheet_path": sheet_path, "is_csv": is_csv, "folder": folder}


def scan_for(p, st, progress):
    sp = st["sheet_path"]
    return scan_folder(st["folder"], p["recursive"], p["skip_hidden"], p["extensions"], p["exclude"],
                       skip_paths=[sp.resolve()] if sp and sp.is_file() else [], progress=progress)


def existing_rows(ws, hr, last, key_col):
    existing = {}
    for r in range(hr + 1, last + 1):
        k = norm_key(ws.cell(row=r, column=key_col).value)
        if k:
            existing.setdefault(k, []).append(r)
    return existing


def make_indexer(ws, hr, last, plan):
    """Running-number columns continue from the highest number already in the sheet."""
    state = {}
    for c, f, prm in plan:
        if f != "index":
            continue
        pat = prm or "#"
        m = re.search(r"#+", pat)
        if not m:
            pat += "#"
            m = re.search(r"#+", pat)
        pre, width, suf = pat[:m.start()], m.end() - m.start(), pat[m.end():]
        rx = re.compile("^" + re.escape(pre) + r"(\d+)" + re.escape(suf) + "$", re.I)
        top = 0
        for r in range(hr + 1, last + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, (int, float)) and not pre and not suf:
                top = max(top, int(v))
            elif v is not None:
                mm = rx.match(str(v).strip())
                if mm:
                    top = max(top, int(mm.group(1)))
        state[c] = [pre, width, suf, top, pat == "#"]

    def next_index(c):
        s = state[c]
        s[3] += 1
        return s[3] if s[4] else f"{s[0]}{str(s[3]).zfill(s[1])}{s[2]}"
    return next_index


def run_job(p, progress=lambda m: None):
    """Scan the folder and update (or create) the spreadsheet. Returns a summary dict."""
    from openpyxl.utils import get_column_letter, range_boundaries

    st = prepare(p, progress)
    wb, meta, ws, hr = st["wb"], st["meta"], st["ws"], st["hr"]
    headers, plan, key_plan = st["headers"], st["plan"], st["key_plan"]
    sheet_path, is_csv = st["sheet_path"], st["is_csv"]

    progress("Scanning folder…")
    files = scan_for(p, st, progress)

    all_cols = sorted(headers)
    template = {}
    if hr + 1 <= ws.max_row:
        template = {c: copy(ws.cell(row=hr + 1, column=c)._style) for c in all_cols}

    last = last_data_row(ws, hr, all_cols)
    if not files and last > hr and p["mode"] != "fresh":
        raise UserError("No files were found in that folder (check the folder and the filters in section 1). "
                        "Nothing was changed.")

    if p["mode"] == "fresh" and last > hr:
        for r in range(hr + 1, last + 1):
            for c in all_cols:
                cell = ws.cell(row=r, column=c)
                if not cell.__class__.__name__ == "MergedCell":
                    cell.value = None
                    cell.hyperlink = None
        last = hr

    existing = existing_rows(ws, hr, last, key_plan[0])
    next_index = make_indexer(ws, hr, last, plan)

    ctx = {"delimiter": p.get("delimiter") or "_", "today": dt.date.today()}
    added = matched = 0
    kc, kf, kp = key_plan
    for i, info in enumerate(files, start=1):
        rows = existing.get(key_text(kf, field_value(kf, kp, info, ctx), is_csv))
        if rows:
            r, is_new = rows.pop(0), False
            matched += 1
        else:
            last += 1
            r, is_new = last, True
            added += 1
            for c, st in template.items():
                ws.cell(row=r, column=c)._style = copy(st)
        for c, f, prm in plan:
            if not is_new:
                if f in NEW_ONLY:
                    continue
                if f != "status" and not p["update_existing"]:
                    continue
            val = next_index(c) if f == "index" else field_value(f, prm, info, ctx)
            put(ws, r, c, f, val, is_csv)
        if i % 500 == 0:
            progress(f"Updating rows… {i:,} of {len(files):,}")

    missing = 0
    status_cols = [(c, prm) for c, f, prm in plan if f == "status"]
    for rows in existing.values():
        for r in rows:
            missing += 1
            for c, prm in status_cols:
                put(ws, r, c, "status", status_texts(prm)[1], is_csv)

    if not is_csv:
        for t in list(ws.tables.values()):
            c1, r1, c2, r2 = range_boundaries(t.ref)
            if r1 == hr and last > r2:
                t.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{last}"
                if t.autoFilter is not None:
                    t.autoFilter.ref = t.ref
        if p["create_new"]:
            ws.auto_filter.ref = f"A1:{get_column_letter(max(all_cols))}{max(last, 1)}"
            for c in all_cols:
                width = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, min(last, 500) + 1))
                ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 10), 60)
        elif ws.auto_filter.ref:
            c1, r1, c2, r2 = range_boundaries(ws.auto_filter.ref)
            if r1 == hr and last > r2:
                ws.auto_filter.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{last}"

    progress("Saving…")
    backup = make_backup(sheet_path) if (not p["create_new"] and p["backup"]) else None
    try:
        save_any(wb, sheet_path, meta)
    except PermissionError:
        raise UserError(f"Couldn't save {sheet_path.name} — it's probably open in Excel.\n"
                        "Close it there and click Update again.")
    return {"path": str(sheet_path), "files": len(files), "added": added, "matched": matched,
            "missing": missing, "backup": str(backup) if backup else None,
            "status_col": bool(status_cols), "refreshed": p["update_existing"]}


def cell_text(f, v):
    """(what the window shows, what gets copied) for one value."""
    if f == "link":
        name, _uri, path = v
        q = lambda x: x.replace('"', '""')  # noqa: E731
        return name, f'=HYPERLINK("{q(path)}","{q(name)}")'
    if v is None:
        return "", ""
    if isinstance(v, dt.datetime):
        t = v.strftime("%Y-%m-%d %H:%M")
    elif isinstance(v, dt.date):
        t = v.isoformat()
    else:
        t = " ".join(str(v).split())
    return t, t


def preview_job(p, progress=lambda m: None):
    """Scan the folder and build rows laid out in the spreadsheet's column order (nothing is saved)."""
    from openpyxl.utils import get_column_letter
    st = prepare(p, progress)
    ws, hr, headers, plan = st["ws"], st["hr"], st["headers"], st["plan"]
    kc, kf, kp = st["key_plan"]
    progress("Scanning folder…")
    files = scan_for(p, st, progress)

    use_sheet = not p["create_new"] and p.get("only_new", True)
    last = last_data_row(ws, hr, sorted(headers)) if use_sheet else hr
    existing = existing_rows(ws, hr, last, kc) if use_sheet else {}
    next_index = make_indexer(ws, hr, last, plan)

    cols = list(range(min(headers), max(headers) + 1))
    columns = [headers.get(c) or f"(column {get_column_letter(c)})" for c in cols]
    fields = {c: (f, prm) for c, f, prm in plan}
    ctx = {"delimiter": p.get("delimiter") or "_", "today": dt.date.today()}
    rows, skipped = [], 0
    for info in files:
        if existing and existing.get(key_text(kf, field_value(kf, kp, info, ctx), False)):
            skipped += 1
            continue
        row = []
        for c in cols:
            if c not in fields:
                row.append(("", ""))
                continue
            f, prm = fields[c]
            val = next_index(c) if f == "index" else field_value(f, prm, info, ctx)
            row.append(cell_text(f, val))
        rows.append(row)
    return {"columns": columns, "rows": rows, "files": len(files), "skipped": skipped,
            "first_header": columns[0], "first_col": get_column_letter(cols[0]),
            "from_sheet": not p["create_new"]}


# --------------------------------------------------------------------------- settings

def load_settings(path=SETTINGS_FILE):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_settings(data, path=SETTINGS_FILE):
    try:
        Path(path).write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    except Exception:
        pass


# --------------------------------------------------------------------------- GUI

class ScrollFrame(ttk.Frame):
    def __init__(self, parent, height=220):
        super().__init__(parent)
        bg = ttk.Style().lookup("TFrame", "background") or None
        self.canvas = tk.Canvas(self, highlightthickness=0, height=height, background=bg)
        self.vsb = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.inner = ttk.Frame(self.canvas)
        self.inner.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.win = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfigure(self.win, width=e.width))
        self.canvas.configure(yscrollcommand=self.vsb.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        self.vsb.pack(side="right", fill="y")
        self.bind("<Enter>", lambda e: self._bind_wheel(True))
        self.bind("<Leave>", lambda e: self._bind_wheel(False))

    def _bind_wheel(self, on):
        for seq in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            if on:
                self.bind_all(seq, self.on_wheel)
            else:
                self.unbind_all(seq)

    def on_wheel(self, e):
        if getattr(e, "num", None) == 4:
            step = -1
        elif getattr(e, "num", None) == 5:
            step = 1
        else:
            d = e.delta
            step = -int(d / 120) if abs(d) >= 120 else (-1 if d > 0 else 1)
        if self.canvas.yview() != (0.0, 1.0):
            self.canvas.yview_scroll(step, "units")
        return "break"


def set_enabled(w, on):
    if isinstance(w, ttk.Combobox):
        w.configure(state="readonly" if on else "disabled")
    else:
        w.configure(state="normal" if on else "disabled")


class ResultsWindow:
    """A table of the scanned files you can copy straight into any spreadsheet."""

    def __init__(self, root, res):
        self.res = res
        self.top = top = tk.Toplevel(root)
        n = len(res["rows"])
        top.title(f"{APP_NAME} — {n:,} file{'s' if n != 1 else ''} ready to copy")
        top.geometry("1050x580")
        top.minsize(600, 320)
        top.columnconfigure(0, weight=1)
        top.rowconfigure(1, weight=1)
        paste = "⌘V" if sys.platform == "darwin" else "Ctrl+V"

        if res["from_sheet"]:
            where = (f"In your spreadsheet, click the first empty cell under “{res['first_header']}” "
                     f"(column {res['first_col']}) and paste with {paste}.")
        else:
            where = f"Click a cell in any spreadsheet and paste with {paste}."
        extra = f"  ({res['skipped']:,} file(s) already in the spreadsheet were left out.)" if res["skipped"] else ""
        ttk.Label(top, text=where + extra, padding=(10, 8, 10, 4), wraplength=1000).grid(row=0, column=0, sticky="w")

        frame = ttk.Frame(top, padding=(10, 0))
        frame.grid(row=1, column=0, sticky="nsew")
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(0, weight=1)
        ids = [f"c{i}" for i in range(len(res["columns"]))]
        self.tree = tree = ttk.Treeview(frame, columns=ids, show="headings", selectmode="extended")
        for i, (cid, head) in enumerate(zip(ids, res["columns"])):
            sample = [len(r[i][0]) for r in res["rows"][:300]]
            width = min(max([len(head)] + sample) * 8 + 16, 420)
            tree.heading(cid, text=head, anchor="w")
            tree.column(cid, width=max(width, 70), stretch=False, anchor="w")
        for r in res["rows"]:
            tree.insert("", "end", values=[d for d, _ in r])
        ys = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        xs = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        tree.grid(row=0, column=0, sticky="nsew")
        ys.grid(row=0, column=1, sticky="ns")
        xs.grid(row=1, column=0, sticky="ew")

        bar = ttk.Frame(top, padding=10)
        bar.grid(row=2, column=0, sticky="ew")
        self.v_headers = tk.BooleanVar(value=not res["from_sheet"])
        ttk.Checkbutton(bar, text="Include the header row", variable=self.v_headers).pack(side="left")
        self.v_msg = tk.StringVar(value="Tip: select some rows to copy just those.")
        ttk.Label(bar, textvariable=self.v_msg, foreground="gray35").pack(side="left", padx=14)
        ttk.Button(bar, text="Close", command=top.destroy).pack(side="right")
        ttk.Button(bar, text="Copy selected rows", command=lambda: self.copy(selected=True)).pack(side="right", padx=6)
        ttk.Button(bar, text="Copy all  ⧉", command=self.copy).pack(side="right")

        for seq in ("<Control-c>", "<Command-c>"):
            top.bind(seq, lambda e: self.copy(selected=bool(tree.selection())))
        for seq in ("<Control-a>", "<Command-a>"):
            top.bind(seq, lambda e: (tree.selection_set(tree.get_children()), "break")[1])
        top.lift()
        top.focus_force()

    def copy(self, selected=False):
        items = self.tree.get_children()
        if selected:
            chosen = set(self.tree.selection())
            if not chosen:
                self.v_msg.set("Select one or more rows first (or use Copy all).")
                return
            idx = [i for i, it in enumerate(items) if it in chosen]
        else:
            idx = range(len(items))
        lines = []
        if self.v_headers.get():
            lines.append("\t".join(self.res["columns"]))
        for i in idx:
            lines.append("\t".join(c for _, c in self.res["rows"][i]))
        self.top.clipboard_clear()
        self.top.clipboard_append("\n".join(lines))
        self.top.update()
        n = len(idx)
        self.v_msg.set(f"Copied {n:,} row{'s' if n != 1 else ''}. Now paste into your spreadsheet.")


class App:
    def __init__(self, root):
        self.root = root
        root.title(APP_NAME)
        root.minsize(780, 680)
        self.q = queue.Queue()
        self.busy = False
        self.map_rows = []
        self.saved_mapping = {}

        self.v_folder = tk.StringVar()
        self.v_recursive = tk.BooleanVar(value=True)
        self.v_hidden = tk.BooleanVar(value=True)
        self.v_exts = tk.StringVar()
        self.v_exclude = tk.StringVar()
        self.v_sheet_mode = tk.StringVar(value="existing")
        self.v_sheet_path = tk.StringVar()
        self.v_sheet_name = tk.StringVar()
        self.v_header_row = tk.StringVar(value="1")
        self.v_new_cols = tk.StringVar(value=DEFAULT_NEW_COLUMNS)
        self.v_delim = tk.StringVar(value="_")
        self.v_key = tk.StringVar()
        self.v_run_mode = tk.StringVar(value="update")
        self.v_update_existing = tk.BooleanVar(value=True)
        self.v_backup = tk.BooleanVar(value=True)
        self.v_only_new = tk.BooleanVar(value=True)
        self.v_status = tk.StringVar(value="Ready.")

        self._build()
        self._apply_settings(load_settings())
        root.protocol("WM_DELETE_WINDOW", self._on_close)
        root.after(100, self._poll)

    # ---- layout
    def _build(self):
        outer = ttk.Frame(self.root, padding=12)
        outer.pack(fill="both", expand=True)
        outer.columnconfigure(0, weight=1)
        outer.rowconfigure(3, weight=1)
        gray = {"foreground": "gray45"}

        # 1 folder
        f1 = ttk.LabelFrame(outer, text=" 1   Folder to scan ", padding=10)
        f1.grid(row=0, column=0, sticky="ew")
        f1.columnconfigure(1, weight=1)
        ttk.Label(f1, text="Folder:").grid(row=0, column=0, sticky="w")
        ttk.Entry(f1, textvariable=self.v_folder).grid(row=0, column=1, sticky="ew", padx=6)
        ttk.Button(f1, text="Choose…", command=self._browse_folder).grid(row=0, column=2, sticky="ew")
        opts = ttk.Frame(f1)
        opts.grid(row=1, column=1, columnspan=2, sticky="w", pady=(4, 2))
        ttk.Checkbutton(opts, text="Include subfolders", variable=self.v_recursive).pack(side="left")
        ttk.Checkbutton(opts, text="Skip hidden & system files", variable=self.v_hidden).pack(side="left", padx=14)
        ttk.Label(f1, text="Only these types:").grid(row=2, column=0, sticky="w")
        ttk.Entry(f1, textvariable=self.v_exts).grid(row=2, column=1, sticky="ew", padx=6, pady=2)
        ttk.Label(f1, text="e.g. pdf, docx (blank = all)", **gray).grid(row=2, column=2, sticky="w")
        ttk.Label(f1, text="Skip names like:").grid(row=3, column=0, sticky="w")
        ttk.Entry(f1, textvariable=self.v_exclude).grid(row=3, column=1, sticky="ew", padx=6, pady=2)
        ttk.Label(f1, text="e.g. *.tmp, Archive, Old*", **gray).grid(row=3, column=2, sticky="w")

        # 2 spreadsheet
        f2 = ttk.LabelFrame(outer, text=" 2   Tracking spreadsheet ", padding=10)
        f2.grid(row=1, column=0, sticky="ew", pady=(10, 0))
        f2.columnconfigure(1, weight=1)
        modes = ttk.Frame(f2)
        modes.grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 4))
        ttk.Radiobutton(modes, text="Use my existing spreadsheet (.xlsx, .xlsm, .csv)", value="existing",
                        variable=self.v_sheet_mode, command=self._apply_sheet_mode).pack(side="left")
        ttk.Radiobutton(modes, text="Create a new one", value="new",
                        variable=self.v_sheet_mode, command=self._apply_sheet_mode).pack(side="left", padx=14)

        ttk.Label(f2, text="File:").grid(row=1, column=0, sticky="w")
        self.e_sheet = ttk.Entry(f2, textvariable=self.v_sheet_path)
        self.e_sheet.grid(row=1, column=1, sticky="ew", padx=6)
        self.b_sheet = ttk.Button(f2, text="Choose…", command=self._browse_sheet)
        self.b_sheet.grid(row=1, column=2, sticky="ew")

        row2 = ttk.Frame(f2)
        row2.grid(row=2, column=1, columnspan=2, sticky="w", padx=6, pady=(4, 0))
        ttk.Label(f2, text="Sheet:").grid(row=2, column=0, sticky="w", pady=(4, 0))
        self.cb_sheet = ttk.Combobox(row2, textvariable=self.v_sheet_name, state="readonly", width=24)
        self.cb_sheet.pack(side="left")
        self.cb_sheet.bind("<<ComboboxSelected>>", lambda e: self._load_columns(detect=True))
        ttk.Label(row2, text="   Headers are on row:").pack(side="left")
        self.sp_header = ttk.Spinbox(row2, from_=1, to=200, width=5, textvariable=self.v_header_row)
        self.sp_header.pack(side="left", padx=4)
        self.b_reload = ttk.Button(row2, text="Reload columns", command=lambda: self._load_columns(detect=False))
        self.b_reload.pack(side="left", padx=6)

        ttk.Label(f2, text="Columns:").grid(row=3, column=0, sticky="w", pady=(6, 0))
        self.e_newcols = ttk.Entry(f2, textvariable=self.v_new_cols)
        self.e_newcols.grid(row=3, column=1, sticky="ew", padx=6, pady=(6, 0))
        self.b_newcols = ttk.Button(f2, text="Use these columns", command=self._apply_new_columns)
        self.b_newcols.grid(row=3, column=2, sticky="ew", pady=(6, 0))
        self.w_existing = [self.e_sheet, self.b_sheet, self.cb_sheet, self.sp_header, self.b_reload]
        self.w_new = [self.e_newcols, self.b_newcols]

        # 3 mapping
        f3 = ttk.LabelFrame(outer, text=" 3   What goes in each column ", padding=10)
        f3.grid(row=3, column=0, sticky="nsew", pady=(10, 0))
        f3.columnconfigure(0, weight=1)
        f3.rowconfigure(1, weight=1)
        ttk.Label(f3, text="Columns set to “— leave as is —” are never changed, so your own status, "
                           "notes and review columns stay safe.", **gray, wraplength=720
                  ).grid(row=0, column=0, sticky="w", pady=(0, 6))
        self.map_frame = ScrollFrame(f3, height=210)
        self.map_frame.grid(row=1, column=0, sticky="nsew")
        bottom = ttk.Frame(f3)
        bottom.grid(row=2, column=0, sticky="w", pady=(8, 0))
        ttk.Label(bottom, text="Match files to existing rows using column:").pack(side="left")
        self.cb_key = ttk.Combobox(bottom, textvariable=self.v_key, state="readonly", width=24)
        self.cb_key.pack(side="left", padx=6)
        ttk.Label(bottom, text="     File-name separator:").pack(side="left")
        ttk.Entry(bottom, textvariable=self.v_delim, width=4).pack(side="left", padx=4)
        ttk.Label(bottom, text="(for “Part of the file name”)", **gray).pack(side="left")

        # 4 run
        f4 = ttk.LabelFrame(outer, text=" 4   Update ", padding=10)
        f4.grid(row=4, column=0, sticky="ew", pady=(10, 0))
        f4.columnconfigure(0, weight=1)
        ttk.Radiobutton(f4, text="Add new files and keep existing rows (recommended)",
                        value="update", variable=self.v_run_mode).grid(row=0, column=0, sticky="w")
        ttk.Radiobutton(f4, text="Start fresh — clear every row under the headers and rebuild the list",
                        value="fresh", variable=self.v_run_mode).grid(row=1, column=0, sticky="w")
        ttk.Checkbutton(f4, text="Refresh details (size, dates, paths…) of files already listed",
                        variable=self.v_update_existing).grid(row=2, column=0, sticky="w", pady=(6, 0))
        ttk.Checkbutton(f4, text="Back up the spreadsheet before saving (into a “_tracker_backups” folder)",
                        variable=self.v_backup).grid(row=3, column=0, sticky="w")
        ttk.Checkbutton(f4, text="In the copy list, show only files that aren't in the spreadsheet yet",
                        variable=self.v_only_new).grid(row=4, column=0, sticky="w")

        btns = ttk.Frame(f4)
        btns.grid(row=5, column=0, sticky="ew", pady=(10, 0))
        ttk.Button(btns, text="Save setup…", command=self._save_setup).pack(side="left")
        ttk.Button(btns, text="Load setup…", command=self._load_setup).pack(side="left", padx=6)
        self.b_run = ttk.Button(btns, text="Update spreadsheet  ▶", command=self._run)
        self.b_run.pack(side="right")
        self.b_show = ttk.Button(btns, text="Show list to copy  ⧉", command=self._show_list)
        self.b_show.pack(side="right", padx=6)
        self.b_open = ttk.Button(btns, text="Open spreadsheet", command=self._open_sheet)
        self.b_open.pack(side="right")

        status = ttk.Frame(outer)
        status.grid(row=5, column=0, sticky="ew", pady=(8, 0))
        status.columnconfigure(1, weight=1)
        self.pb = ttk.Progressbar(status, mode="indeterminate", length=120)
        self.pb.grid(row=0, column=0, padx=(0, 8))
        self.pb.grid_remove()
        ttk.Label(status, textvariable=self.v_status).grid(row=0, column=1, sticky="w")

    # ---- mapping rows
    def _set_columns(self, headers):
        from openpyxl.utils import get_column_letter
        inner = self.map_frame.inner
        for w in inner.winfo_children():
            w.destroy()
        self.map_rows = []
        inner.columnconfigure(3, weight=1)
        bold = ("TkDefaultFont", 0, "bold")
        if not headers:
            ttk.Label(inner, text="No columns yet — choose a spreadsheet above (or enter columns for a new one).",
                      foreground="gray45").grid(row=0, column=0, sticky="w", pady=6)
            self._refresh_keys()
            return
        for c, t in enumerate(("Column", "Fill with", "Setting", "")):
            ttk.Label(inner, text=t, font=bold).grid(row=0, column=c, sticky="w", padx=(0, 8), pady=(0, 4))
        for i, (col, text) in enumerate(headers, start=1):
            saved = self.saved_mapping.get(text)
            if saved is not None and saved[0] in FIELD_BY_KEY:
                field, param = saved[0], saved[1]
            else:
                field = auto_field(text)
                param = FIELD_BY_KEY[field][2] or ""
            fv = tk.StringVar(value=FIELD_BY_KEY[field][1])
            pv = tk.StringVar(value=param or "")
            short = text if len(text) <= 24 else text[:23] + "…"
            ttk.Label(inner, text=f"{get_column_letter(col)}   {short}", width=26).grid(
                row=i, column=0, sticky="w", padx=(0, 8), pady=2)
            cb = ttk.Combobox(inner, textvariable=fv, values=FIELD_LABELS, state="readonly", width=32)
            cb.grid(row=i, column=1, sticky="w", padx=(0, 8), pady=2)
            en = ttk.Entry(inner, textvariable=pv, width=13)
            en.grid(row=i, column=2, sticky="w", pady=2)
            hint = ttk.Label(inner, foreground="gray45")
            hint.grid(row=i, column=3, sticky="w", padx=(6, 0))
            row = {"col": col, "header": text, "fv": fv, "pv": pv, "entry": en, "hint": hint}
            cb.bind("<<ComboboxSelected>>", lambda e, r=row: self._field_changed(r, user=True))
            # Wheel over a dropdown scrolls the list instead of changing the value
            for seq in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
                cb.bind(seq, self.map_frame.on_wheel)
            self.map_rows.append(row)
            self._field_changed(row, user=False)
        self.map_frame.canvas.yview_moveto(0)
        self._refresh_keys()

    def _field_changed(self, row, user):
        key = LABEL_TO_KEY.get(row["fv"].get(), "")
        _, _, default, hint = FIELD_BY_KEY[key]
        if default is None:
            row["pv"].set("")
            row["entry"].configure(state="disabled")
            row["hint"].configure(text="")
        else:
            row["entry"].configure(state="normal")
            if user:
                row["pv"].set(default)
            row["hint"].configure(text=hint)
        if user:
            self._remember_mapping()
            self._refresh_keys()

    def _mapping(self):
        return [{"col": r["col"], "header": r["header"],
                 "field": LABEL_TO_KEY.get(r["fv"].get(), ""), "param": r["pv"].get()}
                for r in self.map_rows]

    def _remember_mapping(self):
        for m in self._mapping():
            self.saved_mapping[m["header"]] = [m["field"], m["param"]]

    def _refresh_keys(self):
        maps = [m for m in self._mapping() if m["field"] not in BAD_KEYS]
        values = [m["header"] for m in maps]
        self.cb_key.configure(values=values)
        if self.v_key.get() in values:
            return
        for k in KEY_PRIORITY:
            hit = next((m["header"] for m in maps if m["field"] == k), None)
            if hit:
                self.v_key.set(hit)
                return
        self.v_key.set(values[0] if values else "")

    # ---- section 2 behaviour
    def _apply_sheet_mode(self):
        existing = self.v_sheet_mode.get() == "existing"
        for w in self.w_existing:
            set_enabled(w, existing)
        for w in self.w_new:
            set_enabled(w, not existing)
        if existing:
            if Path(self.v_sheet_path.get()).is_file():
                self._load_columns(detect=False, quiet=True)
            else:
                self._set_columns([])
        else:
            self._apply_new_columns()

    def _apply_new_columns(self):
        cols = split_list(self.v_new_cols.get())
        self._remember_mapping()
        self._set_columns(dedupe_headers(list(enumerate(cols, start=1))))

    def _load_columns(self, detect, quiet=False):
        path = self.v_sheet_path.get().strip()
        if not path or not Path(path).is_file():
            if not quiet:
                messagebox.showwarning(APP_NAME, "Choose your tracking spreadsheet first.")
            self._set_columns([])
            return
        self._remember_mapping()
        try:
            info = peek(path, self.v_sheet_name.get() or None,
                        None if detect else max(1, to_int(self.v_header_row.get(), 1)))
        except UserError as e:
            messagebox.showerror(APP_NAME, str(e))
            return
        except Exception as e:
            messagebox.showerror(APP_NAME, f"Couldn't read that spreadsheet:\n{e}")
            return
        self.cb_sheet.configure(values=info["sheets"])
        self.v_sheet_name.set(info["sheet"])
        self.v_header_row.set(str(info["header_row"]))
        self._set_columns(info["headers"])
        n = len(info["headers"])
        self.v_status.set(f"Found {n} column{'s' if n != 1 else ''} on row {info['header_row']} of “{info['sheet']}”.")

    def _browse_folder(self):
        d = filedialog.askdirectory(title="Choose the folder to scan", initialdir=self.v_folder.get() or None)
        if d:
            self.v_folder.set(str(Path(d)))

    def _browse_sheet(self):
        start = self.v_sheet_path.get()
        f = filedialog.askopenfilename(
            title="Choose your tracking spreadsheet",
            initialdir=str(Path(start).parent) if start else (self.v_folder.get() or None),
            filetypes=[("Spreadsheets", "*.xlsx *.xlsm *.csv *.tsv"), ("All files", "*.*")])
        if f:
            self.v_sheet_path.set(str(Path(f)))
            self.v_sheet_name.set("")
            self._load_columns(detect=True)

    def _open_sheet(self):
        p = self.v_sheet_path.get()
        if self.v_sheet_mode.get() == "existing" and Path(p).is_file():
            open_path(p)
        else:
            messagebox.showinfo(APP_NAME, "There's no spreadsheet to open yet.")

    # ---- settings
    def _collect_settings(self):
        self._remember_mapping()
        return {
            "folder": self.v_folder.get(), "recursive": self.v_recursive.get(),
            "skip_hidden": self.v_hidden.get(), "extensions": self.v_exts.get(),
            "exclude": self.v_exclude.get(), "sheet_mode": self.v_sheet_mode.get(),
            "sheet_path": self.v_sheet_path.get(), "sheet_name": self.v_sheet_name.get(),
            "header_row": to_int(self.v_header_row.get(), 1), "new_columns": self.v_new_cols.get(),
            "delimiter": self.v_delim.get(), "key_header": self.v_key.get(),
            "run_mode": self.v_run_mode.get(), "update_existing": self.v_update_existing.get(),
            "backup": self.v_backup.get(), "only_new": self.v_only_new.get(), "mapping": self.saved_mapping,
        }

    def _apply_settings(self, s):
        self.saved_mapping = dict(s.get("mapping", {}))
        self.v_folder.set(s.get("folder", ""))
        self.v_recursive.set(s.get("recursive", True))
        self.v_hidden.set(s.get("skip_hidden", True))
        self.v_exts.set(s.get("extensions", ""))
        self.v_exclude.set(s.get("exclude", ""))
        self.v_sheet_mode.set(s.get("sheet_mode", "existing"))
        self.v_sheet_path.set(s.get("sheet_path", ""))
        self.v_sheet_name.set(s.get("sheet_name", ""))
        self.v_header_row.set(str(s.get("header_row", 1)))
        self.v_new_cols.set(s.get("new_columns", DEFAULT_NEW_COLUMNS))
        self.v_delim.set(s.get("delimiter", "_"))
        self.v_key.set(s.get("key_header", ""))
        self.v_run_mode.set(s.get("run_mode", "update"))
        self.v_update_existing.set(s.get("update_existing", True))
        self.v_backup.set(s.get("backup", True))
        self.v_only_new.set(s.get("only_new", True))
        self._apply_sheet_mode()

    def _save_setup(self):
        f = filedialog.asksaveasfilename(title="Save this setup", defaultextension=".json",
                                         initialfile="My file tracker setup.json",
                                         filetypes=[("File Tracker setup", "*.json")])
        if f:
            save_settings(self._collect_settings(), f)
            self.v_status.set(f"Setup saved to {Path(f).name}.")

    def _load_setup(self):
        f = filedialog.askopenfilename(title="Load a saved setup",
                                       filetypes=[("File Tracker setup", "*.json"), ("All files", "*.*")])
        if f:
            s = load_settings(f)
            if not s:
                messagebox.showerror(APP_NAME, "That file isn't a File Tracker setup.")
                return
            self._apply_settings(s)
            self.v_status.set(f"Loaded setup {Path(f).name}.")

    def _on_close(self):
        save_settings(self._collect_settings())
        self.root.destroy()

    # ---- running
    def _run(self):
        if self.busy:
            return
        folder = self.v_folder.get().strip()
        if not folder or not Path(folder).is_dir():
            messagebox.showwarning(APP_NAME, "Choose the folder to scan (section 1).")
            return
        mapping = self._mapping()
        if not any(m["field"] for m in mapping):
            messagebox.showwarning(APP_NAME, "Choose what goes into at least one column (section 3).")
            return
        create_new = self.v_sheet_mode.get() == "new"
        if create_new:
            path = filedialog.asksaveasfilename(
                title="Save the new tracking spreadsheet as", defaultextension=".xlsx",
                initialfile="File Tracker.xlsx", initialdir=str(Path(folder).parent),
                filetypes=[("Excel workbook", "*.xlsx"), ("CSV file", "*.csv")])
            if not path:
                return
        else:
            path = self.v_sheet_path.get().strip()
            if not Path(path).is_file():
                messagebox.showwarning(APP_NAME, "Choose your tracking spreadsheet (section 2).")
                return
            if self.v_run_mode.get() == "fresh" and not messagebox.askyesno(
                    APP_NAME, "Start fresh?\n\nThis clears every row under the headers — including anything "
                              "typed by hand — and rebuilds the list from the folder."
                              + ("\n\nA backup copy is made first." if self.v_backup.get() else
                                 "\n\nBackups are switched OFF.") + "\n\nContinue?"):
                return
        self._start(self._params(folder, path, create_new, mapping), preview=False)

    def _params(self, folder, path, create_new, mapping):
        return {
            "folder": folder, "recursive": self.v_recursive.get(), "skip_hidden": self.v_hidden.get(),
            "extensions": split_list(self.v_exts.get()), "exclude": split_list(self.v_exclude.get()),
            "sheet_path": path, "create_new": create_new,
            "new_columns": [h for _, h in dedupe_headers(list(enumerate(split_list(self.v_new_cols.get()), 1)))],
            "sheet_name": self.v_sheet_name.get(), "header_row": to_int(self.v_header_row.get(), 1),
            "mapping": mapping, "key_header": self.v_key.get(), "delimiter": self.v_delim.get(),
            "mode": "update" if create_new else self.v_run_mode.get(),
            "update_existing": self.v_update_existing.get(), "backup": self.v_backup.get(),
            "only_new": self.v_only_new.get(),
        }

    def _show_list(self):
        if self.busy:
            return
        folder = self.v_folder.get().strip()
        if not folder or not Path(folder).is_dir():
            messagebox.showwarning(APP_NAME, "Choose the folder to scan (section 1).")
            return
        mapping = self._mapping()
        if not any(m["field"] for m in mapping):
            messagebox.showwarning(APP_NAME, "Choose what goes into at least one column (section 3).")
            return
        create_new = self.v_sheet_mode.get() == "new"
        path = "" if create_new else self.v_sheet_path.get().strip()
        if not create_new and not Path(path).is_file():
            messagebox.showwarning(APP_NAME, "Choose your spreadsheet in section 2, or switch to "
                                             "“Create a new one” and type the columns you want.")
            return
        self._start(self._params(folder, path, create_new, mapping), preview=True)

    def _start(self, params, preview):
        self._set_busy(True)
        threading.Thread(target=self._worker, args=(params, preview), daemon=True).start()

    def _worker(self, params, preview=False):
        try:
            if preview:
                self.q.put(("preview", preview_job(params, progress=lambda m: self.q.put(("progress", m)))))
                return
            res = run_job(params, progress=lambda m: self.q.put(("progress", m)))
            res["create_new"] = params["create_new"]
            self.q.put(("done", res))
        except UserError as e:
            self.q.put(("error", str(e)))
        except Exception as e:
            self.q.put(("error", f"Something went wrong:\n{e}\n\n{traceback.format_exc(limit=4)}"))

    def _set_busy(self, on):
        self.busy = on
        self.b_run.configure(state="disabled" if on else "normal")
        self.b_show.configure(state="disabled" if on else "normal")
        if on:
            self.pb.grid()
            self.pb.start(12)
            self.v_status.set("Working…")
        else:
            self.pb.stop()
            self.pb.grid_remove()

    def _poll(self):
        try:
            while True:
                kind, data = self.q.get_nowait()
                if kind == "progress":
                    self.v_status.set(data)
                elif kind == "error":
                    self._set_busy(False)
                    self.v_status.set("Stopped — nothing was saved.")
                    messagebox.showerror(APP_NAME, data)
                elif kind == "done":
                    self._finish(data)
                elif kind == "preview":
                    self._set_busy(False)
                    save_settings(self._collect_settings())
                    n = len(data["rows"])
                    self.v_status.set(f"Listed {n:,} file(s) — nothing was saved.")
                    if n == 0:
                        messagebox.showinfo(APP_NAME, "No new files to list." if data["skipped"] else
                                            "No files were found (check the folder and filters in section 1).")
                    else:
                        ResultsWindow(self.root, data)
        except queue.Empty:
            pass
        self.root.after(100, self._poll)

    def _finish(self, res):
        self._set_busy(False)
        if res["create_new"]:
            self.v_sheet_mode.set("existing")
            self.v_sheet_path.set(res["path"])
            self.v_sheet_name.set("")
            self._remember_mapping()
            self._apply_sheet_mode()
        save_settings(self._collect_settings())
        name = Path(res["path"]).name
        lines = [f"{res['files']:,} files scanned.",
                 f"• {res['added']:,} new file(s) added"]
        if res["matched"]:
            lines.append(f"• {res['matched']:,} already listed" + (" (details refreshed)" if res["refreshed"] else ""))
        if res["missing"]:
            lines.append(f"• {res['missing']:,} listed file(s) weren't found in the folder"
                         + (" (marked in the sheet)" if res["status_col"] else ""))
        if res["backup"]:
            lines.append(f"\nBackup saved in “{BACKUP_DIR}”.")
        self.v_status.set(f"Done — {res['added']:,} added, {res['missing']:,} missing. Saved {name}.")
        if messagebox.askyesno(APP_NAME, f"{name} is updated.\n\n" + "\n".join(lines) + "\n\nOpen it now?"):
            open_path(res["path"])


def main():
    if sys.platform == "win32":
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    root = tk.Tk()
    root.title(APP_NAME)
    try:
        import openpyxl  # noqa: F401
        ok = True
    except ImportError:
        note = ttk.Label(root, text="Setting up (first run only, needs internet)…", padding=30)
        note.pack()
        root.update()
        ok = ensure_openpyxl()
        note.destroy()
    if not ok:
        messagebox.showerror(APP_NAME,
                             "This app needs the free 'openpyxl' package and couldn't install it by itself.\n\n"
                             "Open Terminal (Mac) or Command Prompt (Windows), run:\n\n"
                             "    python3 -m pip install openpyxl\n\n(on Windows: py -m pip install openpyxl)\n\n"
                             "then start File Tracker again.")
        root.destroy()
        return
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
