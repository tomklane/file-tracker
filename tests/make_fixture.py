import sys, shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table
base = Path(sys.argv[1]); shutil.rmtree(base, ignore_errors=True)
proj = base / "Project Alpha"
for f in ["Drawings/Structural/PRJ001_Beam-Layout_RevA.pdf", "Drawings/Structural/PRJ001_Column-Grid_RevB.pdf",
          "Drawings/Arch/PRJ001_Floor-Plan_RevC.dwg", "Specs/PRJ001_Spec-Concrete_Rev0.docx",
          "Correspondence/2026-10-01 Letter to client.pdf", "readme.txt", ".hidden.txt",
          "Archive/old_thing.pdf", "~$lockfile.docx", "Specs/.DS_Store"]:
    p = proj / f; p.parent.mkdir(parents=True, exist_ok=True); p.write_bytes(b"x" * (len(f) * 300))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Register"
wb.create_sheet("Notes")["A1"] = "other sheet stays"
ws["A1"] = "PROJECT ALPHA — DOCUMENT REGISTER"; ws["A1"].font = Font(bold=True, size=14)
for i, h in enumerate(["No.", "Document No", "Title", "Rev", "Discipline", "File Type", "Last Modified", "Link", "Status", "Notes"], 1):
    c = ws.cell(3, i, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
thin = Side(style="thin")
for i, v in enumerate([1, "PRJ001", "Old manual entry", "A", "Structural", "PDF", None, None, "Approved", "keep me"], 1):
    c = ws.cell(4, i, v); c.border = Border(top=thin, bottom=thin, left=thin, right=thin); c.fill = PatternFill("solid", fgColor="FFF2CC")
ws["L4"] = "=COUNTA(B4:B100)"
ws.add_table(Table(displayName="Register", ref="A3:J4"))
wb.save(base / "Tracker.xlsx")
(base / "list.csv").write_text("Ref;File Name;Folder;Size;Date Modified;On Disk;Comments\nDOC-007;ghost.pdf;x;;;;manual\n", encoding="utf-8")
